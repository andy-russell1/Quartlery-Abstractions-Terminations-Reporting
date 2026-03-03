from __future__ import annotations

from copy import copy
from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

MONEY_FMT = '[$GBP] #,##0.00_);[Red]([$GBP] #,##0.00)'
VAT_RATE = 0.2

# Visible bill header cells on the Bill sheet.
PARAM_CELL_MAP: dict[str, str] = {
    "client_name": "B2",
    "currency": "C3",
    "exchange_rate": "C4",
    "billing_frequency": "C6",
    "client_po_number": "C7",
    "payment_terms": "F1",
    "instruction_number": "F2",
    "previous_bill_date": "F3",
    "next_bill_date": "F4",
    "scope": "F5",
    "vat_applied": "F6",
    "person_to_be_billed": "F7",
    "unit_fee": "E11",  # Used by section formulas ($E$11)
}


def copy_range_with_styles(ws: Any, src_range: str, dest_top_left: str) -> None:
    src_start, src_end = src_range.split(":")
    src_min_col = ws[src_start].column
    src_min_row = ws[src_start].row
    src_max_col = ws[src_end].column
    src_max_row = ws[src_end].row

    dest_cell = ws[dest_top_left]
    row_offset = dest_cell.row - src_min_row
    col_offset = dest_cell.column - src_min_col

    for row in range(src_min_row, src_max_row + 1):
        for col in range(src_min_col, src_max_col + 1):
            s = ws.cell(row=row, column=col)
            d = ws.cell(row=row + row_offset, column=col + col_offset)
            d.value = s.value
            if s.has_style:
                d.font = copy(s.font)
                d.fill = copy(s.fill)
                d.border = copy(s.border)
                d.alignment = copy(s.alignment)
                d.number_format = s.number_format
                d.protection = copy(s.protection)
            if s.comment:
                d.comment = copy(s.comment)


def _copy_sheet_to_new_workbook(src_ws: Any, target_title: str) -> Workbook:
    wb_out = Workbook()
    ws_out = wb_out.active
    ws_out.title = target_title

    max_row = src_ws.max_row
    max_col = src_ws.max_column

    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            s = src_ws.cell(row=row, column=col)
            d = ws_out.cell(row=row, column=col)
            d.value = s.value
            if s.has_style:
                d.font = copy(s.font)
                d.fill = copy(s.fill)
                d.border = copy(s.border)
                d.alignment = copy(s.alignment)
                d.number_format = s.number_format
                d.protection = copy(s.protection)
            if s.comment:
                d.comment = copy(s.comment)
            if s.hyperlink:
                d._hyperlink = copy(s.hyperlink)

    for merged in src_ws.merged_cells.ranges:
        ws_out.merge_cells(str(merged))

    for col_letter, dim in src_ws.column_dimensions.items():
        d = ws_out.column_dimensions[col_letter]
        d.width = dim.width
        d.hidden = dim.hidden
        d.outlineLevel = dim.outlineLevel
        d.collapsed = dim.collapsed
        d.bestFit = dim.bestFit

    for row_idx, dim in src_ws.row_dimensions.items():
        d = ws_out.row_dimensions[row_idx]
        d.height = dim.height
        d.hidden = dim.hidden
        d.outlineLevel = dim.outlineLevel
        d.collapsed = dim.collapsed

    ws_out.page_margins = copy(src_ws.page_margins)
    ws_out.page_setup = copy(src_ws.page_setup)
    ws_out.print_options = copy(src_ws.print_options)
    ws_out.freeze_panes = src_ws.freeze_panes

    return wb_out


def _format_date(v: Any, fmt: str = "%d/%m/%Y") -> str:
    if isinstance(v, date):
        return v.strftime(fmt)
    return str(v) if v is not None else ""


def _write_visible_header(ws: Any, params: dict[str, Any], contract_count: int) -> None:
    for key, cell in PARAM_CELL_MAP.items():
        if key in params:
            val = params[key]
            ws[cell] = _format_date(val) if isinstance(val, date) else val

    start = params.get("billing_period_start")
    end = params.get("billing_period_end")
    if isinstance(start, date) and isinstance(end, date):
        ws["C5"] = f"{start.strftime('%d %B %Y')} - {end.strftime('%d %B %Y')}"
        days = (end - start).days + 1
        prorata = max(days, 0) / 365
    else:
        ws["C5"] = ""
        prorata = 0

    unit_fee = float(params.get("unit_fee", 0) or 0)
    ws["E11"] = unit_fee
    ws["E11"].number_format = MONEY_FMT
    ws["C11"] = f"Charged @ {unit_fee:,.2f} GBP per lease / p.a"

    # Keep the upper fee block self-contained on this single sheet.
    ws["D11"] = contract_count
    ws["D12"] = contract_count
    ws["D13"] = contract_count
    ws["F11"] = contract_count * unit_fee * prorata
    ws["G11"] = "=F11*20%"
    ws["H11"] = "=F11+G11"
    ws["F12"] = 0
    ws["G12"] = 0
    ws["H12"] = "=F12+G12"
    ws["F13"] = 0
    ws["G13"] = 0
    ws["H13"] = "=F13+G13"
    ws["F14"] = "=SUM(F11:F13)"
    ws["G14"] = "=SUM(G11:G13)"
    ws["H14"] = "=SUM(H11:H13)"

    for cell in ("F11", "G11", "H11", "F12", "G12", "H12", "F13", "G13", "H13", "F14", "G14", "H14"):
        ws[cell].number_format = MONEY_FMT


def _write_section(
    ws: Any,
    start_row: int,
    data_2col: pd.DataFrame,
    fee_cell_abs: str = "$E$11",
    chargeable: bool = True,
) -> tuple[int, int]:
    if data_2col.empty:
        total_row = start_row
        ws.cell(row=total_row, column=4, value="Total")
        ws.cell(row=total_row, column=6, value=0)
        ws.cell(row=total_row, column=7, value=0)
        ws.cell(row=total_row, column=8, value=0)
        for c in (6, 7, 8):
            ws.cell(row=total_row, column=c).number_format = MONEY_FMT
        return start_row, total_row

    for i, (_, rec) in enumerate(data_2col.iterrows()):
        r = start_row + i
        ws.cell(row=r, column=2, value=rec["system_id"])
        ws.cell(row=r, column=3, value=rec["property_name"])
        ws.cell(row=r, column=4, value=1)
        if chargeable:
            ws.cell(row=r, column=5, value=f"={fee_cell_abs}")
            ws.cell(row=r, column=6, value=f"=D{r}*E{r}")
            ws.cell(row=r, column=7, value=f"=F{r}*{VAT_RATE}")
            ws.cell(row=r, column=8, value=f"=F{r}+G{r}")
        else:
            ws.cell(row=r, column=5, value=0)
            ws.cell(row=r, column=6, value=0)
            ws.cell(row=r, column=7, value=0)
            ws.cell(row=r, column=8, value=0)
        for c in (5, 6, 7, 8):
            ws.cell(row=r, column=c).number_format = MONEY_FMT

    data_end_row = start_row + len(data_2col) - 1
    total_row = data_end_row + 1
    ws.cell(row=total_row, column=4, value="Total")
    if chargeable:
        ws.cell(row=total_row, column=6, value=f"=SUM(F{start_row}:F{data_end_row})")
        ws.cell(row=total_row, column=7, value=f"=SUM(G{start_row}:G{data_end_row})")
        ws.cell(row=total_row, column=8, value=f"=SUM(H{start_row}:H{data_end_row})")
    else:
        ws.cell(row=total_row, column=6, value=0)
        ws.cell(row=total_row, column=7, value=0)
        ws.cell(row=total_row, column=8, value=0)
    for c in (6, 7, 8):
        ws.cell(row=total_row, column=c).number_format = MONEY_FMT

    return data_end_row, total_row


def generate_bill_bytes(
    template_path: str | Path,
    params: dict[str, Any],
    abstractions_2col: pd.DataFrame,
    terminations_2col: pd.DataFrame,
    contract_count: int,
) -> bytes:
    wb = load_workbook(filename=str(template_path))
    if "Bill" not in wb.sheetnames:
        raise ValueError("Template must contain a sheet named 'Bill'.")

    ws = wb["Bill"]
    _write_visible_header(ws, params, contract_count)

    abs_start_row = 18
    abs_data_end, _ = _write_section(ws, abs_start_row, abstractions_2col)

    if len(abstractions_2col) > 0:
        header_row = abs_data_end + 3
    else:
        header_row = abs_start_row

    copy_range_with_styles(ws, "B16:H17", f"B{header_row}")
    ws.cell(row=header_row, column=2, value="New Terminations Completed")

    term_start_row = header_row + 2
    _, term_total_row = _write_section(ws, term_start_row, terminations_2col, chargeable=False)

    trans_header_row = term_total_row + 2
    copy_range_with_styles(ws, "B16:H17", f"B{trans_header_row}")
    ws.cell(row=trans_header_row, column=2, value="Translations Completed")

    translations_blank_rows = 4
    grand_total_row = trans_header_row + 2 + translations_blank_rows
    ws.cell(row=grand_total_row, column=7, value="Grand Total (GBP)")
    ws.cell(row=grand_total_row, column=8, value='=IFERROR($H$14,0)+SUMIF($D:$D,"Total",$H:$H)')
    ws.cell(row=grand_total_row, column=8).number_format = MONEY_FMT

    for col_idx in range(2, 9):
        col = get_column_letter(col_idx)
        if ws.column_dimensions[col].width is None:
            ws.column_dimensions[col].width = 14

    # Output only one worksheet as requested, without template sidecar artifacts.
    wb_out = _copy_sheet_to_new_workbook(ws, target_title="LA ongoing fees")

    out = BytesIO()
    wb_out.save(out)
    out.seek(0)
    return out.getvalue()
